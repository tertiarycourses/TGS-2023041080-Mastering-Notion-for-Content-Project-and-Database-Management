#!/usr/bin/env python3
"""Replace the placeholder TOC with a page-numbered TOC that is ALSO a live Word field.

Two problems with a naive static TOC:
  1. Page numbers are computed from a PDF rendered when the TOC was still a
     one-line placeholder. Once the real entries are written the TOC grows,
     everything after it shifts down, and every number is wrong by the number of
     pages the TOC added. We fix that by ITERATING: render, measure, rewrite,
     re-render, until the numbers stop moving.
  2. A purely static TOC cannot be refreshed in Word (F9) so it rots. We wrap the
     entries in a real `TOC \\o "1-3" \\h \\z \\u` field, with the computed entries
     as the field's cached result — so it prints correctly everywhere AND updates
     in Word.

Usage: python3 inject_toc2.py <docx> [maxlevel]
       (renders its own PDFs via soffice; needs no PDF argument)
"""
import sys, os, re, subprocess, shutil, tempfile
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

SOFFICE = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
GREY = RGBColor(0x33, 0x33, 0x33)
MARKER = "⁣TOCENTRIES⁣"      # invisible sentinel marking our generated block

def norm(s): return re.sub(r'[^a-z0-9]', '', (s or "").lower())

def render_pdf(docx_path, outdir):
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=False, capture_output=True)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    return pdf if os.path.exists(pdf) else None

def collect_headings(doc, maxlevel):
    out = []
    for p in doc.paragraphs:
        sn = p.style.name
        if sn.startswith("Heading"):
            try: lvl = int(sn.split()[-1])
            except ValueError: continue
            if lvl <= maxlevel and p.text.strip():
                out.append((lvl, p.text.strip()))
    return out

# A generated TOC line ends in a dot leader run followed by a page number, e.g.
# "How to Use This Guide.........5". Any PDF page carrying several of those IS a
# TOC page — and its text contains every heading in the document, so searching it
# makes every entry resolve to the TOC itself.
TOC_LINE = re.compile(r"\.{4,}\s*\d+\s*$", re.M)

def toc_pages(raw_pages):
    """Indices of pages that are the rendered TOC (>=3 dot-leader lines, or the
    'TABLE OF CONTENTS' heading)."""
    out = set()
    for i, txt in enumerate(raw_pages):
        # only the real contents page: dot-leader lines, or the literal caps heading near the front
        if len(TOC_LINE.findall(txt or "")) >= 3 or ("TABLE OF CONTENTS" in (txt or "") and i < 5):
            out.add(i)
    return out

def page_map(pdf_path, heads):
    raw = [pg.extract_text() or "" for pg in PdfReader(pdf_path).pages]
    skip = toc_pages(raw)
    pages = [norm(t) for t in raw]
    # Start after the last TOC page: no real heading can live before the contents.
    start = (max(skip) + 1) if skip else 0
    entries, cursor = [], start
    for lvl, text in heads:
        key = norm(text)[:24]; page = None
        for i in range(cursor, len(pages)):
            if i in skip:
                continue
            if key and key in pages[i]:
                page = i + 1; cursor = i; break
        if page is None:                      # fall back: whole document, still skipping the TOC
            for i in range(start, len(pages)):
                if i not in skip and key and key in pages[i]:
                    page = i + 1; break
        entries.append((lvl, text, page or (start + 1)))
    return entries

def find_anchor(doc):
    """Return every paragraph of the previously generated TOC block, or the
    placeholder on the first pass.

    The generated block runs from the MARKER paragraph (which carries the field
    'begin') to the paragraph carrying the field 'end'. Returning ONLY the marker
    paragraph is what caused entries to accumulate on each pass: the old rows were
    never removed, so the second rewrite appended a duplicate set.
    """
    paras = doc.paragraphs
    start = next((i for i, p in enumerate(paras) if MARKER in p.text), None)
    if start is not None:
        end = start
        for j in range(start, len(paras)):
            if 'w:fldCharType="end"' in paras[j]._p.xml:
                end = j
                break
        return paras[start:end + 1], True
    for p in paras:
        xml = p._p.xml
        if ("Update Field" in p.text) or ("TOC " in xml and "instrText" in xml) \
           or ("fldSimple" in xml and "TOC" in xml):
            return [p], False
    return [], False

def _fld(par, kind, instr=None):
    r = par.add_run()._r
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r.append(fc)
    if instr is not None:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = instr
        r2 = par.add_run()._r; r2.append(it)

def write_toc(docx_path, entries, maxlevel):
    doc = Document(docx_path)
    block, is_ours = find_anchor(doc)
    if not block:
        print("  [inject_toc2] no TOC anchor found"); return False
    anchor_el = block[0]._p
    parent = block[0]._parent
    # Remove every other paragraph of the previous block; the first stays only as
    # the insertion point and is deleted at the end.
    for p in block[1:]:
        p._p.getparent().remove(p._p)

    def new_par():
        el = anchor_el.makeelement(qn("w:p"), {})
        anchor_el.addprevious(el)
        return Paragraph(el, parent)

    # --- field begin + instruction, carrying the marker so we can find it again
    p0 = new_par()
    r = p0.add_run(MARKER); r.font.size = Pt(1); r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    _fld(p0, "begin", f'TOC \\o "1-{maxlevel}" \\h \\z \\u')
    sep = p0.add_run()._r
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); sep.append(fc)

    # --- cached result: one paragraph per entry, dot leader + page number
    for lvl, text, page in entries:
        par = new_par()
        pf = par.paragraph_format
        pf.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if lvl >= 2: pf.left_indent = Inches(0.3)
        pf.space_after = Pt(3)
        rr = par.add_run(text + "\t" + str(page))
        rr.font.size = Pt(11 if lvl == 1 else 10.5)
        rr.font.name = "Arial"; rr.bold = (lvl == 1); rr.font.color.rgb = GREY

    # --- field end
    pend = new_par()
    er = pend.add_run()._r
    fce = OxmlElement("w:fldChar"); fce.set(qn("w:fldCharType"), "end"); er.append(fce)
    pend.paragraph_format.space_after = Pt(0)

    anchor_el.getparent().remove(anchor_el)
    doc.save(docx_path)
    return True

def main():
    docx_path = sys.argv[1]
    maxlevel = int(sys.argv[2]) if len(sys.argv) > 2 else 2
    tmp = tempfile.mkdtemp(prefix="toc_")
    prev = None
    for attempt in range(1, 6):
        pdf = render_pdf(docx_path, tmp)
        if not pdf:
            print("  [inject_toc2] render failed"); return
        doc = Document(docx_path)
        heads = collect_headings(doc, maxlevel)
        entries = page_map(pdf, heads)
        nums = [e[2] for e in entries]
        if nums == prev:
            print(f"  [inject_toc2] {os.path.basename(docx_path)}: converged after "
                  f"{attempt-1} rewrite(s) — {len(entries)} entries, pages {nums[0]}..{nums[-1]}")
            return
        prev = nums
        write_toc(docx_path, entries, maxlevel)
    print(f"  [inject_toc2] {os.path.basename(docx_path)}: stopped after 5 passes "
          f"— {len(entries)} entries, pages {nums[0]}..{nums[-1]}")

if __name__ == "__main__":
    main()
