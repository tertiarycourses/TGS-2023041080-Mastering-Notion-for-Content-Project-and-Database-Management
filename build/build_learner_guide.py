#!/usr/bin/env python3
"""Generate the Learner Guide as BOTH a Markdown mirror (LEARNER-GUIDE.md at the repo
root) and a DOCX (courseware/LG-*.docx) from one content stream, so they never diverge.

The LG is where the DETAILED CLICK-BY-CLICK STEPS live — the slide deck carries none.
"""
import os, sys, glob
HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
REPO=os.path.dirname(HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
import prodoc
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4
ASSETS=os.path.join(REPO,"courseware","assets")

B=[]
def h1(t): B.append(("h1",t))
def h2(t): B.append(("h2",t))
def h3(t): B.append(("h3",t))
def p(t):  B.append(("p",t))
def bullets(xs): B.append(("bullets",list(xs)))
def steps(xs): B.append(("steps",list(xs)))
def quote(t): B.append(("quote",t))
def note(t): B.append(("note",t))
def table(headers,rows): B.append(("table",headers,rows))
def rule(): B.append(("rule",))
def figure(path,caption): B.append(("figure",path,caption))
def lab_folder(n):
    g=glob.glob(os.path.join(REPO,"labs",f"lab-{n:02d}-*"))
    return os.path.basename(g[0]) if g else f"lab-{n:02d}"

# ================================================================= CONTENT
h1("How to Use This Guide")
p(f"This Learner Guide accompanies the WSQ course {C.TITLE} ({C.COURSE_CODE}), conducted by {C.ORG}. "
  f"It contains the full click-by-click procedure for all {len(ACT)} hands-on labs, organised by the nine course topics. "
  "The slide deck used in class is deliberately visual and concept-led — every detailed procedure lives here and in the instruction PDF inside each lab folder.")
p(C.SCENARIO)
h2("What you need")
bullets(["A laptop with a modern browser (Chrome, Edge, Safari or Firefox) or the Notion desktop app, and a stable internet connection.",
         "A free Notion account (notion.com → Sign up). A personal email is fine; you will create the Brightwave Media workspace in Lab 1.",
         "The lab folders from the LMS (Activities) or the GitHub repository — one folder per lab with its mock data.",
         "For Lab 12, access to Notion AI: Free and Plus workspaces receive complimentary trial responses; the trainer can add you to the course Business workspace if needed.",
         "For Lab 11, a terminal with curl or Python 3 to call the Notion API (optional but recommended)."])
h2("Course and lab file access")
bullets([f"Course materials (slides, this guide, lesson plan, lab folders) are on the LMS/TMS portal at {C.LMS_URL}.",
         f"The lab folders are also published at {C.GITHUB_REPO} — clone the repository or download the ZIP.",
         "You may use this guide, the slides and your own completed lab work during the open-book assessment."])
h2("How each lab section is organised")
bullets(["Goal and what you will build — the outcome in one paragraph.",
         "Before you begin — the files and pages that must already be in place.",
         "Detailed procedure — numbered steps with the exact menu names, keyboard shortcuts and expected results.",
         "Acceptance check — how you prove the lab is complete.",
         "Evidence to retain — the screenshots to keep for the practical assessment."])

h1("Course Learning Outcomes")
bullets(C.LEARNING_OUTCOMES)

h1("Assessed Knowledge and Abilities")
p("The Written Assessment (9 questions, 1 hour) covers every knowledge statement (K); the Practical Performance (5 tasks, 1 hour) covers every ability statement (A). Each lab states the codes it develops.")
table(["Code","Statement","Developed in"],
      [[k,v,", ".join(f"Lab {a['num']:02d}" for a in ACT if a["km"]==k or a["am"]==k)] for k,v in list(C.KNOWLEDGE)+list(C.ABILITIES)])

h1("Notion Quick Reference")
h3("Keyboard shortcuts and Markdown")
table(["Action","Mac","Windows / Linux","Markdown alternative"],
 [["New page","Cmd + N","Ctrl + N","—"],["Quick Find (search)","Cmd + P or Cmd + K","Ctrl + P or Ctrl + K","—"],
  ["Bold / italic / underline","Cmd + B / I / U","Ctrl + B / I / U","**bold**  *italic*"],
  ["Inline code","Cmd + E","Ctrl + E","`code`"],["Link selected text","Cmd + K","Ctrl + K","—"],
  ["Heading 1 / 2 / 3","Cmd + Option + 1 / 2 / 3","Ctrl + Shift + 1 / 2 / 3","# / ## / ###"],
  ["Bulleted / numbered list","Cmd + Shift + 5 / 6","Ctrl + Shift + 5 / 6","- item  /  1. item"],
  ["To-do / toggle","Cmd + Shift + 4 / 7","Ctrl + Shift + 4 / 7","[] item  /  > item"],
  ["Comment","Cmd + Shift + M","Ctrl + Shift + M","—"],["Highlight (last colour)","Cmd + Shift + H","Ctrl + Shift + H","—"],
  ["Open Notion AI","Space on an empty line; Cmd + Shift + J","Space on an empty line; Ctrl + Shift + J","/ai"],
  ["Any block","/ then the block name","/ then the block name","---  divider, ``` code, \" quote"]])
h3("Menus you will use constantly")
table(["Symbol","Where","What it does"],
 [["⋮⋮","Left of any block on hover","Block menu: Turn into, Duplicate, Delete, Move to, Colour, Copy link; drag to move"],
  ["⋯","Top right of a page; next to a sidebar page","Page menu: font, small text, full width, lock, customise page, page history, export, delete, Turn into wiki"],
  ["Share","Top right of a page","Invite people and groups, set permission levels, Publish to web"],
  ["+ New / New ▾","Top right of a database","Add a row; the ▾ manages templates"],
  ["Filter · Sort · ⋯ (view options)","Above a database","Filters, sorts, groups, properties, layout, lock database, automations"],
  ["@  [[  +","While typing","Mention a person/page/date; link or create a page; create a page elsewhere"]])
h3("Conventions used in this guide")
bullets(["Menu paths are written with arrows, for example Share → Publish → Publish to web.",
         "Text you type is shown in italics or in a fenced block; property names and page titles are in Title Case as they appear in Notion.",
         "'Cmd/Ctrl' means Command on Mac and Control on Windows.",
         "Each lab assumes the previous labs are complete; if you start mid-course, import the previous labs' mock data first.",
         "Screenshots for evidence must not show passwords, personal emails of others or API tokens."])

# ---------------- per-topic, per-lab ----------------
for t in C.TOPICS:
    labs=[a for a in ACT if a["topic"]==t["num"]]
    h1(f"Topic {t['code']} — {t['title']}")
    p(f"{t['lu']} · {t['ka']} · Day {t['day']}. {t['subtitle']}.")
    h3("Key concepts")
    bullets([f"{ttl} — {cap}" for ttl,cap in t["concepts"]])
    for a in labs:
        h2(f"Lab {a['num']:02d} — {a['title']}")
        p(f"Alignment: {a['km']} · {a['am']}.  Level: {a['level']}.  Tools: {a['apps']}.")
        p(f"Objective: {a['objective']}.")
        p(f"Goal: {a['desc']}")
        h3("What you will build")
        p(a["build"])
        fig=os.path.join(ASSETS,"figures",f"lab-{a['num']:02d}-workflow.png")
        if os.path.exists(fig):
            figure(fig, f"Figure {a['num']}. Lab {a['num']:02d} workflow — the four stages and the artifact produced.")
        h3("Before you begin")
        bl=[f"Open the lab folder labs/{lab_folder(a['num'])} and read README.md and lab.yaml.",
            "Have the Brightwave Media workspace open in Notion from the previous labs."]
        for fn,d in a["files"]: bl.append(f"Confirm the file {fn} is present — {d}.")
        bullets(bl)
        h3("Detailed procedure")
        steps(a["steps"])
        h3("Acceptance check")
        p(a["test"])
        h3("Evidence to retain")
        bullets([a["evidence"],"The ticked Steps checklist in the lab README (or lab.yaml status).",
                 "Credential-safe screenshots only — never capture passwords, other people's emails or API tokens."])
        rule()

h1("Assessment")
p("The course is assessed by two open-book instruments on Day 2 from 4:30pm. You may use this guide, the slide deck and your own lab work.")
bullets([C.ASSESSMENT["written"],C.ASSESSMENT["practical"],C.ASSESSMENT["note"],
         "Remember to take the assessment digital attendance (TRAQOM) and complete the TRAQOM survey.",
         f"Submit your answers on the LMS at {C.LMS_URL}.",
         "Oral clarification may be used to close minor gaps; a re-assessment and appeal process is available."])
h3("On assessment day")
steps([("Complete the TRAQOM survey","Scan the QR code shown on the LMS."),
       ("Take the assessment digital attendance","Scan the SSG QR code the trainer displays."),
       ("Sit the Written Assessment (SAQ)","One hour, 9 questions (K1–K9), open book, individual work."),
       ("Sit the Practical Performance (PP)","One hour, 5 tasks (A1–A7) in Notion, open book, individual work; paste screenshots as evidence."),
       ("Submit your answers on the LMS","Check that every question and task has an answer and evidence."),
       ("Sign the Assessment Summary Record","Confirm your result has been recorded with the assessor.")])

h1("Support")
bullets([f"Course portal: {C.LMS_URL}",f"Course page: {C.COURSE_URL}",f"Lab repository: {C.GITHUB_REPO}",
         f"Email: enquiry@tertiaryinfotech.com · Tel: +65 6100 0613",f"Provider: {C.ORG} ({C.UEN})"])

# ================================================================= RENDERERS
def render_md(path):
    L=[f"# {C.TITLE} — Learner Guide\n",
       f"**Course code:** {C.COURSE_CODE}  ·  **Version:** {C.VERSION}  ·  {C.VERSION_DATE}  \n**Conducted by:** {C.ORG} ({C.UEN})\n"]
    for b in B:
        k=b[0]
        if   k=="h1": L.append(f"\n## {b[1]}\n")
        elif k=="h2": L.append(f"\n### {b[1]}\n")
        elif k=="h3": L.append(f"\n**{b[1]}**\n")
        elif k=="p":  L.append(b[1]+"\n")
        elif k=="quote": L.append("> "+b[1]+"\n")
        elif k=="note":  L.append(f"> **Note:** {b[1]}\n")
        elif k=="bullets": L.extend(f"- {x}" for x in b[1]); L.append("")
        elif k=="steps":
            for i,(ttl,d) in enumerate(b[1],1): L.append(f"{i}. **{ttl}**  "); L.append(f"   {d}")
            L.append("")
        elif k=="table":
            hs,rows=b[1],b[2]
            L.append("| "+" | ".join(hs)+" |"); L.append("|"+"|".join(["---"]*len(hs))+"|")
            for r in rows: L.append("| "+" | ".join(str(c).replace("|","\\|") for c in r)+" |")
            L.append("")
        elif k=="figure":
            rel=os.path.relpath(b[1],REPO); L.append(f"![{b[2]}]({rel})\n"); L.append(f"*{b[2]}*\n")
        elif k=="rule": L.append("\n---\n")
    L.append(f"\n---\n\n*{C.TITLE} · {C.COURSE_CODE} · Version {C.VERSION} · © 2026 {C.ORG}*")
    open(path,"w").write("\n".join(L))

def render_docx(path):
    doc=Document()
    st=doc.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(11)
    for s in doc.sections:
        s.top_margin=Inches(0.9); s.bottom_margin=Inches(0.9); s.left_margin=Inches(1.0); s.right_margin=Inches(1.0)
    prodoc.style_headings(doc)
    prodoc.add_cover_page(doc,"Learner Guide",C.TITLE,C.VERSION.lstrip("v")+".0",
        org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),course_code=C.COURSE_CODE)
    prodoc.add_version_control(doc,[
        ("1.0","15 August 2024","First version","Tertiary Infotech Academy"),
        ("2.0","19 March 2026","Second version","Tertiary Infotech Academy"),
        ("6.0",C.VERSION_DATE,f"Full rebuild aligned to the v6 slide deck and Lesson Plan: detailed click-by-click procedures for all {len(ACT)} labs on the Brightwave Media scenario, per-lab mock data folders, Notion quick reference, updated Notion AI, database (data sources) and Notion API content from the Notion Help Center and developer documentation.",C.TRAINER)])
    prodoc.add_toc(doc)
    for b in B:
        k=b[0]
        if   k=="h1": doc.add_heading(b[1],level=1)
        elif k=="h2": doc.add_heading(b[1],level=2)
        elif k=="h3": doc.add_heading(b[1],level=3)
        elif k=="p":  doc.add_paragraph(b[1])
        elif k=="quote":
            pr=doc.add_paragraph(); r=pr.add_run(b[1]); r.italic=True; pr.paragraph_format.left_indent=Inches(0.35)
        elif k=="note":
            pr=doc.add_paragraph(); r=pr.add_run("Note: "+b[1]); r.italic=True; r.font.color.rgb=prodoc.GREY
        elif k=="bullets":
            for x in b[1]: doc.add_paragraph(str(x),style="List Bullet")
        elif k=="steps":
            for i,(ttl,d) in enumerate(b[1],1):
                pr=doc.add_paragraph(); r=pr.add_run(f"Step {i}. {ttl}"); r.bold=True
                pr.paragraph_format.space_before=Pt(8); pr.paragraph_format.space_after=Pt(2)
                doc.add_paragraph(d)
        elif k=="table":
            hs,rows=b[1],b[2]
            t=doc.add_table(rows=0,cols=len(hs)); t.style="Table Grid"
            hr=t.add_row().cells
            for i,h in enumerate(hs):
                hr[i].text=""; r=hr[i].paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
                prodoc._shade_cell(hr[i],"1F6FEB")
            for row in rows:
                c=t.add_row().cells
                for i,v in enumerate(row):
                    c[i].text=""; rr=c[i].paragraphs[0].add_run(str(v)); rr.font.size=Pt(9.5)
            doc.add_paragraph()
        elif k=="figure":
            ip=doc.add_paragraph(); ip.alignment=WD_ALIGN_PARAGRAPH.CENTER; ip.add_run().add_picture(b[1],width=Inches(6.3))
            cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cr=cp.add_run(b[2]); cr.font.size=Pt(9); cr.italic=True; cr.font.color.rgb=prodoc.GREY
        elif k=="rule":
            pr=doc.add_paragraph(); pr.paragraph_format.space_after=Pt(2)
            r=pr.add_run("_"*88); r.font.color.rgb=RGBColor(0xDD,0xE3,0xEC); r.font.size=Pt(8)
    prodoc.add_page_numbers(doc); prodoc.enable_update_fields(doc)
    doc.save(path)

if __name__=="__main__":
    md=os.path.join(REPO,"LEARNER-GUIDE.md"); dx=os.path.join(REPO,"courseware",f"LG-{C.SHORT_TITLE}.docx")
    render_md(md); render_docx(dx)
    print("Saved",md); print("Saved",dx); print(f"blocks {len(B)}, labs {len(ACT)}")
