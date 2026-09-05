#!/usr/bin/env python3
"""Generate the Lesson Plan (LP) DOCX in the WSQ house format.

Two days, 9:30am–6:30pm with a 1-hour lunch. Day 1 carries 480 instructional
minutes; Day 2 carries 360 instructional minutes plus the 120-minute assessment
(WA 60 + PP 60 from 4:30pm) — 14 training hours + 2 assessment hours in total,
matching the TGS record. Slide references are read from slide_map.json written by
build_slides.py, so they can never drift from the deck.
"""
import os, sys, json
HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
REPO=os.path.dirname(HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
import prodoc
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4
ASSETS=os.path.join(REPO,"courseware","assets")
BY={a["num"]:a for a in ACT}

SM=json.load(open(os.path.join(HERE,"slide_map.json")))
ORDER=sorted(SM.items(),key=lambda kv:kv[1])
def rng(key):
    """Slide range from a mark to the slide before the next mark."""
    start=SM[key]; later=[v for k,v in ORDER if v>start]
    end=(min(later)-1) if later else start
    return f"Slide {start}" if end<=start else f"Slides {start}–{end}"
def topic_rng(n):
    """Concept slides of a topic: from the topic mark to the first lab of that topic."""
    start=SM[f"topic{n}"]; labs=[a["num"] for a in ACT if a["topic"]==n]
    end=(SM[f"lab{labs[0]}"]-1) if labs else start
    return f"Slides {start}–{end}"
def lab_rng(n): return rng(f"lab{n}")
def T(n): return [t for t in C.TOPICS if t["num"]==n][0]

def _clock(m): h,mm=divmod(m,60); return f"{h}:{mm:02d}"
def _span(t0,d): return f"{_clock(t0)}–{_clock(t0+d)}", t0+d

def lab_row(n,mins):
    a=BY[n]
    return (mins,"lab",f"Lab {n:02d}: {a['title']} ({a['km']} · {a['am']})","Demonstration, then individual hands-on practice with acceptance check",lab_rng(n))

def build_rows(items):
    rows=[]; t0=9*60+30
    for mins,kind,text,method,ref in items:
        if kind=="lunch":
            sp=f"{_clock(t0)}–{_clock(t0+60)}"; t0+=60
            rows.append((sp,"—",text,"Break","—")); continue
        sp,t0=_span(t0,mins); rows.append((sp,f"{mins} min",text,method,ref))
    return rows,t0

DAY1=[
 (20,"admin","Welcome, AM digital attendance (TRAQOM), introductions, ground rules, learning outcomes and course outline","Facilitated briefing",f"Slides {SM['admin']}–{SM['topic1']-1}"),
 (45,"topic",f"Topic 01: {T(1)['title']} — {T(1)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(1)),
 lab_row(1,30),
 (10,"break","Tea break","Break","—"),
 lab_row(2,35),
 (40,"topic",f"Topic 02: {T(2)['title']} — {T(2)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(2)),
 lab_row(3,35),
 (0,"lunch","Lunch (1 hour, not counted in the instructional hours)","",""),
 (10,"admin","PM digital attendance (TRAQOM)","Admin","—"),
 (50,"topic",f"Topic 03: {T(3)['title']} — {T(3)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(3)),
 lab_row(4,40),
 lab_row(5,40),
 (10,"break","Tea break","Break","—"),
 (40,"topic",f"Topic 04: {T(4)['title']} — {T(4)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(4)),
 lab_row(6,45),
 (30,"recap","Day 1 recap, Q&A and preview of Day 2","Facilitated review",rng("day1end")),
]
DAY2=[
 (15,"admin","Day 1 recap and AM digital attendance (TRAQOM)","Facilitated briefing","—"),
 (40,"topic",f"Topic 05: {T(5)['title']} — {T(5)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(5)),
 lab_row(7,30),
 (10,"break","Tea break","Break","—"),
 lab_row(8,30),
 (35,"topic",f"Topic 06: {T(6)['title']} — {T(6)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(6)),
 lab_row(9,40),
 (0,"lunch","Lunch (1 hour, not counted in the instructional hours)","",""),
 (5,"admin","PM digital attendance (TRAQOM)","Admin","—"),
 (25,"topic",f"Topic 07: {T(7)['title']} — {T(7)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(7)),
 lab_row(10,25),
 (25,"topic",f"Topic 08: {T(8)['title']} — {T(8)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(8)),
 lab_row(11,25),
 (20,"topic",f"Topic 09: {T(9)['title']} — {T(9)['subtitle']}","Lecture with visual models, demonstration, didactic questioning",topic_rng(9)),
 lab_row(12,25),
 (10,"admin","Course summary, briefing for assessment, assessment digital attendance (TRAQOM) and TRAQOM survey","Facilitated briefing",f"Slides {SM['close']}–{SM['end']-1}"),
 (60,"assess","Written Assessment (WA) — Short-Answer Questions, 9 questions (K1–K9), open book","Individual assessment","—"),
 (60,"assess","Practical Performance (PP) — 5 scenario-based tasks in Notion (A1–A7), open book","Individual assessment","—"),
]

def minutes(items,kinds): return sum(m for m,k,*_ in items if k in kinds)

def build():
    doc=Document()
    st=doc.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(11)
    for s in doc.sections:
        s.top_margin=Inches(0.9); s.bottom_margin=Inches(0.9); s.left_margin=Inches(1.0); s.right_margin=Inches(1.0)
    prodoc.style_headings(doc)
    prodoc.add_cover_page(doc,"LESSON PLAN",C.TITLE,C.VERSION.lstrip("v")+".0",
        org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),course_code=C.COURSE_CODE)
    prodoc.add_version_control(doc,[
        ("1.0","7 March 2024","First version","Tertiary Infotech Academy"),
        ("2.0","16 October 2024","Second version","Tertiary Infotech Academy"),
        ("3.0","6 August 2025","New format","Tertiary Infotech Academy"),
        ("4.0","17 October 2025","Update company name","Tertiary Infotech Academy"),
        ("6.0",C.VERSION_DATE,"Full rebuild aligned to the v6 slide deck (single-source pipeline): twelve hands-on labs on the Brightwave Media scenario with mock data, exact slide references per session, Day 2 assessment block 4:30–6:30pm, Notion AI and Notion API content updated from the Notion Help Center and developer documentation.",C.TRAINER)])
    prodoc.add_toc(doc)

    doc.add_heading("Course Overview",level=1)
    doc.add_paragraph(f"{C.TITLE} is a two-day, {C.TRAINING_HOURS}-hour WSQ programme (plus {C.ASSESSMENT_HOURS} hours of assessment) mapped to the Skills Framework TSC {C.TSC_TITLE} ({C.TSC_CODE}). "
        "Day 1 builds the foundations: the workspace, pages and blocks, a governed wiki with policies and verified pages, collaboration tools, and rich, styled content. "
        "Day 2 turns the workspace into a content management system: databases with views, relations, rollups, formulas and automations; sharing, permissions and web publishing with page analytics; imports, integrations and the Notion API; and Notion AI. "
        f"{C.SCENARIO}")
    t=doc.add_table(rows=0,cols=2); t.style="Table Grid"
    for k,v in [("Course Code",C.COURSE_CODE),("Skills Framework TSC",f"{C.TSC_TITLE} · {C.TSC_CODE}"),
                ("Training Duration",f"{C.TRAINING_HOURS} hours across {C.DAYS} days (9:30am–6:30pm, 1-hour lunch)"),
                ("Assessment","WA (SAQ) 60 minutes + PP 60 minutes on Day 2 from 4:30pm — both open book"),
                ("Mode of Delivery","Classroom (or synchronous online), instructor-led, with individual hands-on labs in Notion"),
                ("Trainer",C.TRAINER),("Class Size","Minimum 3, maximum 20 learners"),
                ("Attendance","A minimum of 75% attendance is required for assessment and funding")]:
        c=t.add_row().cells
        c[0].text=""; r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(10)
        c[1].text=""; r2=c[1].paragraphs[0].add_run(v); r2.font.size=Pt(10)
    doc.add_paragraph()

    doc.add_heading("Learning Outcomes",level=1)
    for lo in C.LEARNING_OUTCOMES: doc.add_paragraph(lo,style="List Bullet")

    doc.add_heading("Learning Units, Topics and Labs",level=1)
    t=doc.add_table(rows=0,cols=4); t.style="Table Grid"
    hr=t.add_row().cells
    for i,h in enumerate(["Learning Unit","Topics","Labs","LO · K/A"]):
        hr[i].text=""; r=hr[i].paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
        prodoc._shade_cell(hr[i],"1F6FEB")
    for code,name,topics,lo,ka in C.LEARNING_UNITS:
        labs=[a for a in ACT if a["topic"] in topics]
        c=t.add_row().cells
        for i,v in enumerate([f"{code}: {name}","; ".join(f"Topic {n}: {T(n)['title']}" for n in topics),
                              ", ".join(f"Lab {a['num']:02d}" for a in labs),f"{lo} · {ka}"]):
            c[i].text=""; rr=c[i].paragraphs[0].add_run(v); rr.font.size=Pt(9.5)
    doc.add_paragraph()

    doc.add_heading("Assessed Knowledge and Abilities",level=1)
    doc.add_paragraph("The Written Assessment covers every knowledge statement (K1–K9) and the Practical Performance covers every ability statement (A1–A7).")
    t=doc.add_table(rows=0,cols=3); t.style="Table Grid"
    hr=t.add_row().cells
    for i,h in enumerate(["Code","Statement","Developed in"]):
        hr[i].text=""; r=hr[i].paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
        prodoc._shade_cell(hr[i],"1F6FEB")
    for code,txt in list(C.KNOWLEDGE)+list(C.ABILITIES):
        labs=[a for a in ACT if a["km"]==code or a["am"]==code]
        c=t.add_row().cells
        for i,v in enumerate([code,txt,", ".join(f"Lab {a['num']:02d}" for a in labs)]):
            c[i].text=""; rr=c[i].paragraphs[0].add_run(v); rr.font.size=Pt(9.5); rr.bold=(i==0)
    doc.add_paragraph()

    doc.add_heading("Daily Schedule",level=1)
    doc.add_paragraph("Slide references point to the current deck, "+f"{C.SHORT_TITLE}-{C.VERSION}.pptx. Tea breaks are counted within training time; lunch is not.")
    for day,items in ((1,DAY1),(2,DAY2)):
        rows,tend=build_rows(items)
        train=minutes(items,("admin","topic","lab","break","recap")); assess=minutes(items,("assess",))
        doc.add_heading(f"Day {day}: {C.DAY_THEMES[day]}",level=2)
        doc.add_paragraph(f"Instructional time {train} minutes ({train//60} hours)" + (f"; assessment {assess} minutes ({assess//60} hours)" if assess else "") + f"; class ends {_clock(tend)}.")
        t=doc.add_table(rows=0,cols=5); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        hr=t.add_row().cells
        for i,h in enumerate(["Time","Duration","Topic / Activity","Method","Slides"]):
            hr[i].text=""; r=hr[i].paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
            prodoc._shade_cell(hr[i],"1F6FEB")
        for row in rows:
            c=t.add_row().cells
            for i,v in enumerate(row):
                c[i].text=""; rr=c[i].paragraphs[0].add_run(str(v)); rr.font.size=Pt(9)
        for row in t.rows:
            row.cells[0].width=Inches(0.95); row.cells[1].width=Inches(0.7); row.cells[2].width=Inches(3.0); row.cells[3].width=Inches(1.3); row.cells[4].width=Inches(0.85)
        doc.add_paragraph()
        assert (train==480 if day==1 else (train==360 and assess==120)), (day,train,assess)

    doc.add_heading("Topic-by-Topic Delivery Guide",level=1)
    for tp in C.TOPICS:
        labs=[a for a in ACT if a["topic"]==tp["num"]]
        doc.add_heading(f"Topic {tp['code']}: {tp['title']} · {topic_rng(tp['num'])}",level=2)
        doc.add_paragraph(f"{tp['lu']} · {tp['ka']} · Day {tp['day']}. {tp['subtitle']}.")
        doc.add_paragraph("Key concepts delivered:")
        for ttl,cap in tp["concepts"]: doc.add_paragraph(f"{ttl} — {cap}",style="List Bullet")
        doc.add_paragraph("Facilitator guidance: demonstrate each concept live in the trainer's Notion workspace, invite learners to share how they handle the same task today, and use the lab immediately after so the concept is practised while fresh.")
        for a in labs:
            doc.add_heading(f"Lab {a['num']:02d}: {a['title']} · {lab_rng(a['num'])}",level=3)
            doc.add_paragraph(f"Alignment: {a['km']} · {a['am']}. Level: {a['level']}. Tools: {a['apps']}. Duration: see schedule.")
            doc.add_paragraph(f"Learner output: {a['build']}.")
            doc.add_paragraph(f"Acceptance check: {a['test']}")
            doc.add_paragraph(f"Lab folder: labs/lab-{a['num']:02d}-… containing " + ", ".join(fn for fn,_ in a["files"]) + f", LAB-{a['num']:02d}-Instructions.pdf and README.md.")

    doc.add_heading("Assessment Plan",level=1)
    for line in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],
                 "Both instruments are open book: learners may use the slides, the Learner Guide and their own lab work.",
                 "The assessment is conducted on Day 2 from 4:30pm to 6:30pm after the assessment digital attendance and the TRAQOM survey.",
                 C.ASSESSMENT["note"],
                 "Oral clarification (up to 5 minutes, 1:1) may be used to close minor gaps. Learners assessed Not Yet Competent may be re-assessed; an appeal process is available."]:
        doc.add_paragraph(line,style="List Bullet")

    doc.add_heading("Resources Required",level=1)
    for r_ in ["A laptop per learner (Windows or Mac) with a modern browser or the Notion desktop app.",
               "A free Notion account per learner; the trainer's Business-plan workspace for Notion AI features when trial responses are exhausted.",
               "The course lab folders (labs/), one folder per lab with mock data, distributed through the LMS and GitHub.",
               "Projector or large display, whiteboard and markers for demonstrations.",
               "The slide deck, Learner Guide and Lesson Plan available on the LMS.",
               "SSG digital attendance QR codes for AM, PM and assessment sessions."]:
        doc.add_paragraph(r_,style="List Bullet")

    prodoc.add_page_numbers(doc); prodoc.enable_update_fields(doc)
    out=os.path.join(REPO,"courseware",f"LP-{C.SHORT_TITLE}.docx")
    doc.save(out); print("Saved",out)

if __name__=="__main__": build()
